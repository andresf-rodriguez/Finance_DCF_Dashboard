"""
Dashboard de valoracion. Ejecutar con:  streamlit run app.py
Lee datos_{TICKER}.csv (creado por extraer_datos.py). Si no existe, lo genera.
El Excel del DCF se construye en memoria a partir de plantilla_DCF.xlsx al pulsar "Descargar Excel".
"""
import io
import json
import math
import os
import re
import subprocess
import sys
from datetime import date

import pandas as pd
import plotly.graph_objects as go
import streamlit as st
from openpyxl import load_workbook

st.set_page_config(page_title="Valoración de empresas", page_icon="📊", layout="wide")

# Rutas absolutas respecto a este archivo: la app no depende de la carpeta desde la que se lance streamlit
BASE = os.path.dirname(os.path.abspath(__file__))
EXTRACTOR = os.path.join(BASE, "extraer_datos.py")
PLANTILLA = os.path.join(BASE, "plantilla_DCF.xlsx")


def ruta_datos(ticker: str, ext: str) -> str:
    """datos_{ticker}.csv / .json junto a app.py (donde los escribe extraer_datos.py con cwd=BASE)."""
    return os.path.join(BASE, f"datos_{ticker}.{ext}")


# ---------- Utilidades ----------
METRICAS = {
    "Flujo operativo (OCF)": "CFO",
    "Flujo de caja libre (FCF)": "FCF",
    "FCF menos compensación en acciones": "FCF - SBC",
}
EXPLICA = {
    "CFO": "La caja que genera el negocio ANTES de inversiones (capex). Asume que toda la inversión es opcional: "
           "es un TOPE OPTIMISTA. En empresas que invierten mucho (IA, centros de datos) sobrevalora.",
    "FCF": "Caja después de inversiones. Si la empresa está en un ciclo de inversión fuerte, castiga de más: "
           "es un PISO CONSERVADOR.",
    "FCF - SBC": "Igual que FCF, pero restando lo que la empresa paga a empleados en acciones (un costo real que diluye "
                 "al accionista). La medida más exigente.",
}


def entorno_sec():
    """Entorno para lanzar extraer_datos.py: le pasa el User-Agent de la SEC desde st.secrets
    (si existe) como variable SEC_USER_AGENT. Si falta, el script avisa como configurarlo."""
    env = dict(os.environ)
    try:
        if "SEC_USER_AGENT" in st.secrets:
            env["SEC_USER_AGENT"] = st.secrets["SEC_USER_AGENT"]
    except Exception:
        pass  # sin ningun archivo de secrets, st.secrets lanza excepcion; el script avisara como configurarlo
    return env


FILAS_EXCEL = {  # fila de la hoja "Datos" de la plantilla para cada concepto
    "Ingresos": 6, "CFO": 7, "Capex": 8, "Arrendamientos financieros": 9, "SBC": 10,
    "Caja": 11, "Valores negociables": 12, "Deuda": 13, "Acciones diluidas": 14,
}


def excel_dcf(ticker: str, df: pd.DataFrame, rellenados: list, precio: float) -> bytes:
    """Construye en memoria el Excel del DCF: copia de plantilla_DCF.xlsx con la hoja "Datos"
    rellenada con los ultimos 5 años y el precio actual en Supuestos!B6. Las hojas "DCF" y
    "Supuestos" son formulas que leen de "Datos", asi que no hay que tocarlas."""
    ultimos = df.iloc[:, -5:]
    wb = load_workbook(PLANTILLA)
    datos = wb["Datos"]
    datos["B1"] = ticker
    datos["B2"] = date.today().isoformat()
    datos["C2"] = ("El precio de Supuestos!B6 es el de esta fecha de descarga. Crecimiento, beta, prima, coste de deuda, "
                   "impuestos, g y margen de seguridad son valores genéricos: revísalos antes de sacar conclusiones.")
    for j, anio in enumerate(ultimos.columns):
        datos.cell(row=5, column=2 + j, value=int(anio))
        for dato, fila in FILAS_EXCEL.items():
            valor = ultimos.at[dato, anio]
            if dato == "Acciones diluidas" and anio in rellenados:
                valor = float("nan")  # en la app se rellena con el año vecino; en el Excel va vacio
            # Asignar .value directamente: cell(value=None) no borra el ejemplo que trae la plantilla
            datos.cell(row=fila, column=2 + j).value = None if pd.isna(valor) else float(valor)
    wb["Supuestos"]["B6"] = float(precio)
    wb.calculation.fullCalcOnLoad = True  # Excel recalcula todo al abrir
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@st.cache_data(ttl=3600)
def cargar(ticker: str):
    """Devuelve (df, meta). meta viene de datos_{ticker}.json (fecha del último 10-K,
    años aproximados, reexpresiones). Si el JSON no existe, meta queda vacío."""
    ruta = ruta_datos(ticker, "csv")
    if not os.path.exists(ruta):
        subprocess.run([sys.executable, EXTRACTOR, ticker], check=True, capture_output=True, env=entorno_sec(), cwd=BASE)
    df = pd.read_csv(ruta, index_col=0)
    df.columns = [int(c) for c in df.columns]
    meta = {}
    if os.path.exists(ruta_datos(ticker, "json")):
        with open(ruta_datos(ticker, "json")) as f:
            meta = json.load(f)
    # Años sin acciones diluidas: usar el año válido más cercano (y avisar después)
    acc = df.loc["Acciones diluidas"]
    meta["rellenados"] = [int(a) for a in acc.index[acc.isna()]]
    df.loc["Acciones diluidas"] = acc.ffill().bfill()
    return df, meta


def variantes(ticker: str) -> list:
    """Formas del ticker a probar en Yahoo Finance. Las acciones con clase se escriben BRK.B en bolsa,
    pero Yahoo y la SEC usan BRK-B: se prueba primero con guion para no hacer una peticion fallida."""
    return list(dict.fromkeys([ticker.replace(".", "-"), ticker, ticker.replace("-", ".")]))


@st.cache_data(ttl=3600)
def splits_posteriores(ticker: str, desde: str):
    """Splits de Yahoo Finance con fecha posterior a `desde` (presentación del último 10-K).
    Si el split fue antes de presentar el 10-K, el 10-K ya viene reexpresado y no hay que tocar nada."""
    try:
        import yfinance as yf
    except Exception:
        return []
    for t in variantes(ticker):
        try:
            s = yf.Ticker(t).splits
            if s.index.tz is not None:
                s.index = s.index.tz_convert(None)
            s = s[s.index > pd.Timestamp(desde)]
            return [(d.strftime("%d-%m-%Y"), float(r)) for d, r in s.items() if r > 0]
        except Exception:
            continue
    return []

@st.cache_data(ttl=900)
def precio_actual(ticker: str):
    try:
        import yfinance as yf
    except Exception:
        return None
    for t in variantes(ticker):
        try:
            p = yf.Ticker(t).fast_info["last_price"]
            if p is not None and float(p) > 0:
                return float(p)
        except Exception:
            continue
    return None

def fmt_b(x): return f"\\${x:,.1f} B"
def fmt_p(x): return f"\\${x:,.2f}"


def punto_de_partida(etiqueta, valor_10k, clave, anio):
    """Campo editable con el valor del 10-K por defecto. Devuelve (valor, editado).
    Si el usuario lo cambia, muestra el valor original y un botón para restaurarlo."""
    st.session_state.setdefault(clave, valor_10k)
    valor = st.number_input(etiqueta, key=clave, step=0.1, format="%.3f",
                            help="Por defecto, el valor del último 10-K. Puedes escribir otro "
                                 "(por ejemplo, una cifra más reciente que conozcas).")
    editado = not math.isclose(valor, valor_10k, rel_tol=1e-9, abs_tol=1e-9)
    if editado:
        st.caption(f"✏️ Editado. El 10-K de {anio} reporta {fmt_b(valor_10k)}.")
        st.button("Restaurar valor del 10-K", key=f"restaurar_{clave}",
                  on_click=lambda: st.session_state.update({clave: valor_10k}))
    return float(valor), editado


# ---------- Aviso (primer elemento de la página; se muestra en todos los casos, tambien si la carga falla) ----------
st.warning(
    "**Herramienta educativa. No es asesoramiento financiero ni una recomendación de compra o venta.**\n\n"
    "- Los márgenes salen del histórico de la empresa, pero **el crecimiento y la tasa de descuento (WACC, g) son "
    "valores genéricos**: revísalos antes de sacar conclusiones.\n"
    "- Un DCF es **muy sensible al WACC y al crecimiento perpetuo (g)**: un punto de diferencia cambia mucho el valor.\n"
    "- Datos: SEC EDGAR (10-K). Precio: Yahoo Finance, del momento de la consulta. Verifica en la fuente original."
)

# ---------- Barra lateral: empresa ----------
st.sidebar.header("Empresa")
ticker = st.sidebar.text_input("Ticker (bolsa de EE. UU.)", "META").strip().upper()
# El ticker va a nombres de archivo y a un subprocess: validar antes de usarlo en cualquier sitio
if not re.fullmatch(r"[A-Z0-9.\-]{1,6}", ticker):
    st.error("Ticker no válido. Usa solo letras, números, punto o guion, máximo 6 caracteres (ej.: META, BRK.B).")
    st.stop()
if st.sidebar.button("🔄 Actualizar datos desde la SEC"):
    with st.spinner("Descargando de la SEC..."):
        r = subprocess.run([sys.executable, EXTRACTOR, ticker], capture_output=True, text=True, env=entorno_sec(), cwd=BASE)
    st.cache_data.clear()
    for k in [k for k in st.session_state if k.startswith(("base_", "ing_base_"))]:
        del st.session_state[k]  # los puntos de partida editados vuelven al nuevo 10-K
    # Sentencias, no una expresion suelta: Streamlit ("magic") pasaria el DeltaGenerator resultante
    # a st.write y acabaria mostrando la documentacion completa de Streamlit en la pagina.
    if r.returncode == 0:
        st.sidebar.success("Listo")
    else:
        st.sidebar.error(r.stderr[-800:])

try:
    df, meta = cargar(ticker)
except Exception as e:
    print(f"[cargar] {ticker}: {type(e).__name__}: {e}", file=sys.stderr)  # detalle solo en la consola del servidor
    st.error(f"No pude cargar {ticker}. Comprueba que el ticker existe en la bolsa de EE. UU. y que la SEC "
             "tiene datos de esa empresa (formulario 10-K).")
    st.stop()

anios = list(df.columns)
ultimo = anios[-1]

# ---------- Revisar la serie de acciones (splits, años faltantes, aproximados) ----------
avisos = []
if meta.get("filed"):
    splits = splits_posteriores(ticker, meta["filed"])
    factor_split = 1.0
    for _, r in splits:
        factor_split *= r
    if factor_split != 1.0:
        df.loc["Acciones diluidas"] *= factor_split
        detalle = ", ".join(f"{r:g}:1 el {d}" for d, r in splits)
        avisos.append(("info", f"Acciones ajustadas por split ({detalle}) posterior al último 10-K "
                               f"(presentado el {meta['filed']}), para que coincidan con el precio actual."))
else:
    avisos.append(("warning", "Estos datos se generaron con una versión anterior del extractor. Pulsa "
                              "«Actualizar datos desde la SEC» para revisar splits y acciones faltantes."))
if meta.get("rellenados"):
    avisos.append(("warning", f"No hay acciones diluidas para {', '.join(map(str, meta['rellenados']))}; "
                              "se usó el año válido más cercano."))
aprox = sorted({a["anio"] for a in meta.get("aproximados", []) if a["dato"] == "Acciones diluidas"})
if aprox:
    avisos.append(("info", f"Acciones diluidas de {', '.join(map(str, aprox))} estimadas como utilidad neta ÷ EPS "
                           "diluido, porque la SEC no publica el total consolidado. Marcadas con * en la tabla."))
if (df.loc["Acciones diluidas"].pct_change().abs() > 0.5).any():
    avisos.append(("warning", "El número de acciones cambia más de 50 % entre años: puede ser un split sin ajustar. "
                              "Pulsa «Actualizar datos desde la SEC»."))
if df.loc["Acciones diluidas"].isna().all():
    st.error(f"No encontré el número de acciones de {ticker} en ningún año. No puedo valorar por acción.")
    st.stop()
precio_auto = precio_actual(ticker)
precio = st.sidebar.number_input("Precio actual de la acción ($)", min_value=0.01,
                                 value=float(precio_auto) if precio_auto else 100.0, step=1.0,
                                 help="Se intenta traer de Yahoo Finance. Si no aparece, escríbelo a mano.")
if not precio_auto:
    st.sidebar.caption("No pude obtener el precio automáticamente (instala yfinance: `pip install yfinance`).")

# ---------- Descargar Excel (se construye en memoria solo al pulsar) ----------
if not os.path.exists(PLANTILLA):
    st.sidebar.caption(f"No encuentro {os.path.basename(PLANTILLA)} junto a app.py, así que no se puede descargar el Excel del DCF.")
else:
    pocos_anios = df.shape[1] < 5
    st.sidebar.download_button(
        "📥 Descargar Excel del DCF",
        data=lambda: excel_dcf(ticker, df, meta.get("rellenados", []), precio),
        file_name=f"DCF_{ticker}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        disabled=pocos_anios,
        help="Copia de plantilla_DCF.xlsx con los últimos 5 años de esta empresa y el precio actual.",
    )
    if pocos_anios:
        st.sidebar.caption("La plantilla necesita 5 años de datos y esta empresa tiene menos.")

acciones = df.at["Acciones diluidas", ultimo]
caja = df.at["Caja", ultimo] + df.at["Valores negociables", ultimo]
deuda = df.at["Deuda", ultimo]
cap_mercado = precio * acciones

# ---------- Encabezado ----------
st.title(f"📊 {ticker} — ¿Cuánto vale la acción?")
c1, c2, c3, c4 = st.columns(4)
c1.metric("Precio actual", fmt_p(precio))
c2.metric("Capitalización", fmt_b(cap_mercado))
c3.metric("Precio / OCF", f"{cap_mercado / df.at['CFO', ultimo]:.1f}x")
fcf = df.at["FCF", ultimo]
c4.metric("Precio / FCF", f"{cap_mercado / fcf:.1f}x" if fcf > 0 else "n/a")
st.caption(f"Datos de la SEC hasta el año fiscal {ultimo}. Cifras en miles de millones de USD.")
if meta.get("cierre") and meta.get("filed"):
    cierre = pd.Timestamp(meta["cierre"])
    meses = (date.today().year - cierre.year) * 12 + date.today().month - cierre.month
    st.info(f"📅 Los datos corresponden al último año fiscal completo reportado en el 10-K: cierre el "
            f"{cierre:%d-%m-%Y}, presentado el {pd.Timestamp(meta['filed']):%d-%m-%Y}. No incluyen trimestres "
            f"posteriores; han pasado {meses} meses desde el cierre. Si la empresa cambió desde entonces "
            "(ventas, márgenes, deuda, acciones), el resultado puede variar.")
else:
    st.info(f"📅 Los datos corresponden al último año fiscal completo reportado en el 10-K ({ultimo}). No incluyen "
            "trimestres posteriores. Si la empresa cambió desde entonces, el resultado puede variar.")
for tipo, msg in avisos:
    (st.info if tipo == "info" else st.warning)(msg)

with st.expander("Ver datos históricos"):
    if aprox:
        tabla = df.apply(lambda col: col.map(lambda v: "" if pd.isna(v) else f"{v:,.2f}"))
        for a in aprox:
            tabla.at["Acciones diluidas", a] += "*"
        st.dataframe(tabla, width="stretch")
        st.caption("\\* Estimado como utilidad neta ÷ EPS diluido.")
    else:
        st.dataframe(df.round(2), width="stretch")

tab_simple, tab_dcf = st.tabs(["🟢 Modo simple", "🔵 Modo DCF (detallado)"])

# =====================================================================
# MODO SIMPLE: proyecta una métrica, aplica un múltiplo, descuenta
# =====================================================================
with tab_simple:
    izq, der = st.columns([1, 2])
    with izq:
        nombre = st.selectbox("¿Con qué medir la caja que genera?", list(METRICAS))
        met = METRICAS[nombre]
        st.info(EXPLICA[met])
        serie = df.loc[met]
        base_10k = float(serie[ultimo])
        base, base_editada = punto_de_partida(f"Punto de partida: {nombre} (miles de millones USD)",
                                              base_10k, f"base_{ticker}_{met}", ultimo)
        # El crecimiento histórico sale siempre de la serie del 10-K, no del valor editado
        cagr5 = (serie.iloc[-1] / serie.iloc[0]) ** (1 / (len(serie) - 1)) - 1 if serie.iloc[0] > 0 and base_10k > 0 else None
        n = st.slider("Años a proyectar", 3, 10, 5)
        g0 = st.slider("Crecimiento anual inicial (%)", -10.0, 40.0,
                       float(round(min(max(cagr5 * 100, -10), 40), 1)) if cagr5 else 10.0, 0.5,
                       help=f"Crecimiento histórico compuesto de los últimos {len(serie)-1} años: "
                            f"{cagr5*100:.1f}%" if cagr5 else "Histórico no calculable (valores negativos).")
        decay = st.slider("Desaceleración del crecimiento por año (%)", 0, 30, 10,
                          help="0 = el crecimiento se mantiene. 10 = cada año crece 10% menos rápido que el anterior.")
        multiplo = st.slider("Múltiplo de salida (precio / métrica)", 5.0, 40.0, 15.0, 0.5,
                             help="A cuántas veces la métrica se venderá la empresa al final. Mira el múltiplo actual arriba.")
        retorno = st.slider("Retorno anual que exiges (%)", 5.0, 20.0, 10.0, 0.5)
        acc_valida = df.loc["Acciones diluidas"].dropna()
        acc0, acc1 = acc_valida.iloc[0], acc_valida.iloc[-1]
        n_acc = acc_valida.index[-1] - acc_valida.index[0]
        d_hist = ((acc1 / acc0) ** (1 / n_acc) - 1) * 100 if acc0 > 0 and acc1 > 0 and n_acc > 0 else 0.0
        d_hist = float(min(max(round(d_hist, 2), -5), 5))
        d_acc = st.slider("Cambio anual en acciones (%)", -5.0, 5.0, d_hist, 0.25,
                          help="Negativo = recompras (menos acciones, más valor por acción). Si el histórico falta, arranca en 0.")

    # Proyeccion
    proy, tasas, g, val = [], [], g0 / 100, base
    for _ in range(n):
        val *= 1 + g
        proy.append(val)
        tasas.append(g)
        g *= 1 - decay / 100
    acc_fin = acciones * (1 + d_acc / 100) ** n
    precio_futuro = proy[-1] * multiplo / acc_fin
    # Con flujo <= 0 no hay múltiplo que aplicar (y una raíz de un negativo da un complejo)
    aplica = base > 0 and precio_futuro > 0
    if aplica:
        valor_justo = precio_futuro / (1 + retorno / 100) ** n
        cagr_esperado = (precio_futuro / precio) ** (1 / n) - 1
        dif = valor_justo / precio - 1

    with der:
        if aplica:
            color = "green" if dif > 0 else "red"
            st.markdown(f"### Valor justo: :{color}[{fmt_p(valor_justo)}]  &nbsp;&nbsp; "
                        f"({dif:+.1%} vs. precio actual)")
            st.caption("⚠️ Depende del múltiplo de salida, el crecimiento y el retorno exigido que has puesto a la "
                       "izquierda; son valores de ejemplo hasta que los ajustes.")
            st.markdown(f"Comprando hoy a {fmt_p(precio)} y vendiendo en {n} años a {fmt_p(precio_futuro)}, "
                        f"el retorno anual sería **{cagr_esperado:.1%}** (tú exiges {retorno:.1f}%).")
        else:
            st.warning("Con caja negativa o cero este modelo no aplica: no se puede proyectar un múltiplo sobre "
                       "un flujo ≤ 0. Usa OCF como métrica o el modo DCF.")
        if met == "CFO":
            capex_pct = (df.at["Capex", ultimo] + df.at["Arrendamientos financieros", ultimo]) / df.at["CFO", ultimo]
            if capex_pct > 0.4:
                st.warning(f"⚠️ {ticker} reinvierte el {capex_pct:.0%} de su flujo operativo. Con OCF estás asumiendo "
                           "que toda esa inversión podría dejar de hacerse sin afectar el negocio. Compara con FCF.")
        fig = go.Figure()
        fig.add_bar(x=[str(a) for a in anios], y=serie.values, name="Histórico", marker_color="#2e8b57")
        if base_editada:
            fig.add_bar(x=["Punto de partida"], y=[base], name="Tu punto de partida", marker_color="#ff7f0e")
        fig.add_bar(x=[str(ultimo + i + 1) for i in range(n)], y=proy, name="Proyección", marker_color="#1f77b4")
        fig.update_layout(title=f"{nombre} (miles de millones USD)", height=420, barmode="group",
                          legend=dict(orientation="h", y=-0.15))
        # Eje categórico: si no, Plotly trata los años como números y descarta la categoría de texto
        orden_x = [str(a) for a in anios] + (["Punto de partida"] if base_editada else []) + \
                  [str(ultimo + i + 1) for i in range(n)]
        fig.update_xaxes(type="category", categoryorder="array", categoryarray=orden_x)
        st.plotly_chart(fig, width="stretch")
        with st.expander("🧮 ¿Cómo se calcula?"):
            origen = (f"**editado por ti; el 10-K de {ultimo} reporta {fmt_b(base_10k)}**" if base_editada
                      else f"año fiscal {ultimo}")
            st.markdown(f"**Paso 1: proyectar {nombre}.** Punto de partida: {fmt_b(base)} ({origen}). "
                        f"Crecimiento inicial {g0:.1f}%; cada año se multiplica por (1 − {decay}%).")
            for i, (t, v) in enumerate(zip(tasas, proy), 1):
                st.markdown(f"- Año {i} ({ultimo + i}): crece {t:.2%} → {fmt_b(v)}")
            st.markdown(f"**Paso 2: acciones en {ultimo + n}.** {acciones:.3f} B × (1 + ({d_acc:.2f}%))^{n} = **{acc_fin:.3f} B**")
            if aplica:
                st.markdown(f"**Paso 3: precio futuro.** {fmt_b(proy[-1])} × {multiplo}x ÷ {acc_fin:.3f} B acciones = **{fmt_p(precio_futuro)}**")
                st.markdown(f"**Paso 4: valor justo hoy.** {fmt_p(precio_futuro)} ÷ (1 + {retorno:.1f}%)^{n} = **{fmt_p(valor_justo)}**")
                st.markdown(f"**Paso 5: retorno esperado.** ({fmt_p(precio_futuro)} ÷ {fmt_p(precio)})^(1/{n}) − 1 = **{cagr_esperado:.1%}**")
            else:
                st.markdown("**Pasos 3 a 5: no aplican.** El flujo proyectado es ≤ 0, así que no hay múltiplo ni precio futuro que calcular.")

# =====================================================================
# MODO DCF
# =====================================================================
with tab_dcf:
    st.markdown("Aquí el capex no se ignora: **tú decides cuándo y hasta dónde baja**. "
                "Todo se proyecta como % de ingresos.")
    ing = df.loc["Ingresos"]
    cfo_m = (df.loc["CFO"] / ing)
    capex_m = ((df.loc["Capex"] + df.loc["Arrendamientos financieros"]) / ing)
    sbc_m = (df.loc["SBC"] / ing)
    g_hist = (ing.iloc[-1] / ing.iloc[0]) ** (1 / (len(ing) - 1)) - 1

    a, b, c = st.columns(3)
    with a:
        st.markdown("**Ingresos**")
        st.caption(f"Histórico: {g_hist:.1%} anual compuesto; último año {ing.iloc[-1]/ing.iloc[-2]-1:.1%}")
        ing_base, ing_editada = punto_de_partida("Ingresos base (miles de millones USD)", float(ing.iloc[-1]),
                                                 f"ing_base_{ticker}", ultimo)
        g_ini = st.slider("Crecimiento año 1 (%)", -5.0, 40.0, 12.0, 0.5)
        g_fin = st.slider("Crecimiento año 10 (%)", 0.0, 15.0, 4.0, 0.5)
        st.markdown("**Margen de flujo operativo**")
        st.caption(f"Histórico: último {cfo_m.iloc[-1]:.1%}, promedio 3 años {cfo_m.iloc[-3:].mean():.1%}")
        m_cfo = st.slider("CFO como % de ingresos", 5.0, 80.0, float(round(cfo_m.iloc[-3:].mean() * 100, 1)), 0.5)
    with b:
        st.markdown("**Inversión (capex + arrendamientos)**")
        st.caption(f"Histórico: último {capex_m.iloc[-1]:.1%}, promedio 3 años {capex_m.iloc[-3:].mean():.1%}")
        cx_ini = st.slider("Capex año 1 (% ingresos)", 0.0, 80.0, float(round(capex_m.iloc[-1] * 100, 1)), 0.5)
        cx_fin = st.slider("Capex año 10 (% ingresos)", 0.0, 50.0, float(round(capex_m.iloc[:-1].mean() * 100, 1)), 0.5,
                           help="Nivel 'normal' cuando termine el ciclo de inversión.")
        anio_norm = st.slider("Año en que el capex llega al nivel normal", 2, 10, 5)
        st.markdown("**Compensación en acciones**")
        st.caption(f"Histórico: {sbc_m.iloc[-1]:.1%} de ingresos")
        m_sbc = st.slider("SBC como % de ingresos", 0.0, 25.0, float(round(sbc_m.iloc[-1] * 100, 1)), 0.5)
    with c:
        st.markdown("**Tasa de descuento (WACC)**")
        rf = st.number_input("Tasa libre de riesgo (%)", 0.0, 10.0, 4.5, 0.1, help="Bono del Tesoro EE. UU. a 10 años")
        beta = st.number_input("Beta", 0.2, 3.0, 1.2, 0.05)
        prima = st.number_input("Prima de riesgo de mercado (%)", 2.0, 10.0, 5.0, 0.25)
        kd = st.number_input("Costo de la deuda (%)", 0.0, 15.0, 5.0, 0.25)
        tax = st.number_input("Tasa de impuestos (%)", 0.0, 40.0, 21.0, 0.5)
        ke = rf + beta * prima
        w_e = cap_mercado / (cap_mercado + deuda) if cap_mercado + deuda > 0 else 1
        wacc = w_e * ke + (1 - w_e) * kd * (1 - tax / 100)
        st.metric("WACC", f"{wacc:.2f}%", help=f"Costo del capital {ke:.2f}% × {w_e:.0%} + deuda {kd:.1f}% × (1-imp.) × {1-w_e:.0%}")
        g_perp = st.slider("Crecimiento perpetuo (%)", 0.0, 4.0, 2.5, 0.25)
        mos = st.slider("Margen de seguridad (%)", 0, 50, 20)

    if wacc / 100 <= g_perp / 100:
        st.error("El WACC debe ser mayor que el crecimiento perpetuo.")
        st.stop()

    # Proyeccion 10 años
    filas = []
    rev = ing_base
    for t in range(1, 11):
        g = g_ini + (g_fin - g_ini) * (t - 1) / 9
        rev *= 1 + g / 100
        cx = cx_fin if t >= anio_norm else cx_ini + (cx_fin - cx_ini) * (t - 1) / (anio_norm - 1)
        cfo = rev * m_cfo / 100
        capex = rev * cx / 100
        sbc = rev * m_sbc / 100
        fcf_t = cfo - capex - sbc
        disc = 1 / (1 + wacc / 100) ** t
        filas.append(dict(Año=ultimo + t, Ingresos=rev, CFO=cfo, Capex=capex, SBC=sbc,
                          **{"FCF - SBC": fcf_t, "VP": fcf_t * disc, "_crec": g, "_cx": cx}))
    p = pd.DataFrame(filas)
    suma_vp = p["VP"].sum()
    tv = p["FCF - SBC"].iloc[-1] * (1 + g_perp / 100) / (wacc / 100 - g_perp / 100)
    vp_tv = tv / (1 + wacc / 100) ** 10
    ev = suma_vp + vp_tv
    equity = ev + caja - deuda
    intrinseco = equity / acciones
    con_mos = intrinseco * (1 - mos / 100)

    st.divider()
    r1, r2, r3, r4 = st.columns(4)
    r1.metric("Valor intrínseco por acción", fmt_p(intrinseco), f"{intrinseco/precio-1:+.1%} vs. precio")
    r2.metric("Precio con margen de seguridad", fmt_p(con_mos))
    r3.metric("Valor de la empresa", fmt_b(ev))
    r4.metric("% que viene del valor terminal", f"{vp_tv/ev:.0%}")
    st.caption("⚠️ Este valor sale de los supuestos de arriba, que en parte son valores genéricos (crecimiento, WACC, g) "
               "hasta que los revises tú. No lo uses sin haberlos revisado.")
    if vp_tv / ev > 0.75:
        st.warning("Más del 75% del valor depende del valor terminal: el resultado es muy sensible al WACC y al crecimiento perpetuo.")
    if precio <= con_mos:
        st.success(f"El precio actual ({fmt_p(precio)}) está por debajo del precio con margen de seguridad.")
    elif precio <= intrinseco:
        st.info("El precio actual está entre el precio con margen de seguridad y el valor intrínseco.")
    else:
        st.error("El precio actual está por encima del valor intrínseco según tus supuestos.")

    g1, g2 = st.columns(2)
    with g1:
        fig = go.Figure()
        hist = df.loc["FCF - SBC"]
        fig.add_bar(x=[str(a) for a in anios], y=hist.values, name="Histórico", marker_color="#2e8b57")
        if ing_editada:
            # FCF - SBC implícito por tus ingresos editados y tus márgenes (año 0 de la proyección)
            fig.add_bar(x=["Punto de partida"], y=[ing_base * (m_cfo - cx_ini - m_sbc) / 100],
                        name="Tu punto de partida", marker_color="#ff7f0e")
        fig.add_bar(x=p["Año"].astype(str), y=p["FCF - SBC"], name="Proyección", marker_color="#1f77b4")
        fig.update_layout(title="FCF menos SBC (miles de millones USD)", height=380, legend=dict(orientation="h", y=-0.2))
        orden_x = [str(a) for a in anios] + (["Punto de partida"] if ing_editada else []) + list(p["Año"].astype(str))
        fig.update_xaxes(type="category", categoryorder="array", categoryarray=orden_x)
        st.plotly_chart(fig, width="stretch")
    with g2:
        fig2 = go.Figure()
        fig2.add_scatter(x=[str(a) for a in anios], y=capex_m.values * 100, name="Capex % ingresos (hist.)", line=dict(color="#2e8b57"))
        fig2.add_scatter(x=p["Año"].astype(str), y=p["_cx"], name="Capex % ingresos (tu supuesto)", line=dict(color="#1f77b4", dash="dash"))
        fig2.add_scatter(x=[str(a) for a in anios], y=cfo_m.values * 100, name="CFO % ingresos (hist.)", line=dict(color="#888"))
        fig2.update_layout(title="Márgenes como % de ingresos", height=380, legend=dict(orientation="h", y=-0.2))
        st.plotly_chart(fig2, width="stretch")

    with st.expander("Ver proyección año por año"):
        st.dataframe(p.drop(columns=["_crec", "_cx"]).set_index("Año").round(1), width="stretch")
    with st.expander("🧮 ¿Cómo se calcula el DCF?"):
        st.markdown(f"""
**Paso 1: proyectar ingresos.** Parten de {fmt_b(ing_base)} ({f"**editado por ti; el 10-K de {ultimo} reporta {fmt_b(ing.iloc[-1])}**" if ing_editada else ultimo}). El crecimiento baja en línea recta de {g_ini:.1f}% (año 1) a {g_fin:.1f}% (año 10). Ingresos en el año 10: {fmt_b(p['Ingresos'].iloc[-1])}.

**Paso 2: convertir ingresos en caja.** Cada año: FCF − SBC = Ingresos × (CFO {m_cfo:.1f}% − capex − SBC {m_sbc:.1f}%). El capex baja de {cx_ini:.1f}% a {cx_fin:.1f}% de los ingresos en el año {anio_norm}. FCF − SBC en el año 10: {fmt_b(p['FCF - SBC'].iloc[-1])}.

**Paso 3: tasa de descuento (WACC).** Costo del capital = {rf:.1f}% + {beta:.2f} × {prima:.1f}% = {ke:.2f}%. WACC = {w_e:.0%} × {ke:.2f}% + {1 - w_e:.0%} × {kd:.1f}% × (1 − {tax:.0f}%) = **{wacc:.2f}%**.

**Paso 4: traer los 10 años a hoy.** Cada flujo se divide entre (1 + {wacc:.2f}%)^año. Suma: **{fmt_b(suma_vp)}**.

**Paso 5: valor terminal** (lo que vale la empresa después del año 10). {fmt_b(p['FCF - SBC'].iloc[-1])} × (1 + {g_perp:.1f}%) ÷ ({wacc:.2f}% − {g_perp:.1f}%) = {fmt_b(tv)}. Traído a hoy: **{fmt_b(vp_tv)}** ({vp_tv / ev:.0%} del valor total).

**Paso 6: de empresa a acción.** {fmt_b(suma_vp)} + {fmt_b(vp_tv)} + caja {fmt_b(caja)} − deuda {fmt_b(deuda)} = {fmt_b(equity)} ÷ {acciones:.2f} B acciones = **{fmt_p(intrinseco)}**. Con {mos}% de margen de seguridad: **{fmt_p(con_mos)}**.
""")

    st.markdown("**Sensibilidad del valor intrínseco** (filas: WACC, columnas: crecimiento perpetuo)")
    filas_s = []
    for dw in (-2, -1, 0, 1, 2):
        w = (wacc + dw) / 100
        fila = {}
        for dg in (-1, -0.5, 0, 0.5, 1):
            gg = (g_perp + dg) / 100
            if w <= gg:
                fila[f"g {gg:.1%}"] = "n/a"; continue
            pv = sum(p["FCF - SBC"].iloc[t - 1] / (1 + w) ** t for t in range(1, 11))
            tvs = p["FCF - SBC"].iloc[-1] * (1 + gg) / (w - gg) / (1 + w) ** 10
            fila[f"g {gg:.1%}"] = f"${(pv + tvs + caja - deuda) / acciones:,.0f}"
        filas_s.append(pd.Series(fila, name=f"WACC {w:.1%}"))
    st.table(pd.DataFrame(filas_s))

st.sidebar.divider()
st.sidebar.caption("Datos: SEC EDGAR (10-K). Precio: Yahoo Finance. Esto no es una recomendación de inversión.")